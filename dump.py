"""Dump every tab of the two Blaze Camp workbooks into agent-readable text files.

For each sheet writes dumps/<book>/<NN>_<sheet>.txt containing:
  - merged ranges, data validations (dropdown sources), hyperlinks, comments
  - every non-empty cell, row-major, with fill colour / bold / font colour
and a column-major view (dumps/<book>/<NN>_<sheet>.columns.txt) that lists
each column top-to-bottom - the natural way to read the exercise libraries.
"""
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
BOOKS = {"BYOP": ROOT / "raw/BYOP.xlsx", "MASTER_MODEL": ROOT / "raw/MASTER_MODEL.xlsx"}


def colour(c):
    try:
        if c is None:
            return None
        if c.type == "rgb" and c.rgb and c.rgb not in ("00000000", "FFFFFFFF", "FF000000"):
            return c.rgb[-6:]
        if c.type == "theme":
            return f"theme{c.theme}{'+%.2f' % c.tint if c.tint else ''}"
    except Exception:
        return None
    return None


def cell_attrs(cell):
    bits = []
    fill = cell.fill
    if fill is not None and fill.fill_type == "solid":
        fc = colour(fill.fgColor)
        if fc:
            bits.append(f"bg={fc}")
    if cell.font is not None:
        if cell.font.bold:
            bits.append("bold")
        fc = colour(cell.font.color)
        if fc:
            bits.append(f"fg={fc}")
    return " ".join(bits)


def safe(name):
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")


def main():
    for book, path in BOOKS.items():
        wb = openpyxl.load_workbook(path, data_only=False)
        out = ROOT / "dumps" / book
        out.mkdir(parents=True, exist_ok=True)
        index = []
        for i, ws in enumerate(wb.worksheets):
            fname = f"{i:02d}_{safe(ws.title)}"
            lines = [f"SHEET: {ws.title!r}", f"DIMENSIONS: {ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}"]
            if ws.freeze_panes:
                lines.append(f"FREEZE: {ws.freeze_panes}")
            merged = [str(r) for r in ws.merged_cells.ranges]
            lines.append(f"MERGED ({len(merged)}): " + ", ".join(merged))
            lines.append("DATA VALIDATIONS (dropdown sources):")
            for dv in ws.data_validations.dataValidation:
                lines.append(f"  {dv.sqref} -> type={dv.type} formula1={dv.formula1!r}")
            links = [(c.coordinate, c.hyperlink.target or c.hyperlink.location) for row in ws.iter_rows() for c in row if c.hyperlink]
            if links:
                lines.append("HYPERLINKS:")
                lines += [f"  {a} -> {t}" for a, t in links]
            comments = [(c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment]
            if comments:
                lines.append("COMMENTS:")
                lines += [f"  {a}: {t!r}" for a, t in comments]
            lines.append("")
            lines.append("CELLS (row-major; coord | value | formatting):")
            n = 0
            cols = {}
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None or (isinstance(v, str) and v.strip() == ""):
                        continue
                    n += 1
                    attrs = cell_attrs(c)
                    lines.append(f"{c.coordinate} | {v!r}" + (f" | {attrs}" if attrs else ""))
                    cols.setdefault(c.column, []).append((c.row, v, attrs))
            lines.insert(2, f"NON-EMPTY CELLS: {n}")
            (out / f"{fname}.txt").write_text("\n".join(lines), encoding="utf-8")
            clines = [f"SHEET: {ws.title!r} - COLUMN-MAJOR VIEW (each column top to bottom)"]
            for ci in sorted(cols):
                clines.append("")
                clines.append(f"=== COLUMN {get_column_letter(ci)} ({len(cols[ci])} cells) ===")
                for r, v, a in cols[ci]:
                    clines.append(f"  r{r}: {v!r}" + (f"  [{a}]" if a else ""))
            (out / f"{fname}.columns.txt").write_text("\n".join(clines), encoding="utf-8")
            index.append(f"{fname}.txt\t{ws.title!r}\t{n} cells\t{ws.dimensions}")
        (out / "INDEX.txt").write_text("\n".join(index), encoding="utf-8")
        print(book, "\n  " + "\n  ".join(index))


if __name__ == "__main__":
    sys.exit(main())
