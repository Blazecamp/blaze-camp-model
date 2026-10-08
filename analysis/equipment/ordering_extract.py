"""Extract every session (ordered list of exercise slots) from the program tabs of
BYOP (00-08) and MASTER_MODEL (00,01,02,04,05,06,07,11,12).

Read-only: opens raw/*.xlsx with openpyxl (data_only=True). Never touches Google Sheets.
Each session = {sid, book, tab, track, week, day, level, kind, items:[{cell,row,slot,text}]}
`slot` is the coach's own row label (CATEGORY / PURPOSE / Part label) when one exists, verbatim.
"""
import datetime
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "raw"

_wb_cache = {}


def wb(book):
    if book not in _wb_cache:
        _wb_cache[book] = openpyxl.load_workbook(RAW / f"{book}.xlsx", data_only=True)
    return _wb_cache[book]


def val(ws, r, c):
    if isinstance(c, str):
        c = CI(c)
    x = ws.cell(r, c).value
    if isinstance(x, datetime.datetime):  # 'm/d' formatted reps like 8/8 (see prog_mm_gym.md 0.1)
        x = f"{x.month}/{x.day}"
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    if isinstance(x, str):
        x = x.strip()
        if x == "":
            x = None
    return x


PLACEHOLDER = {"CHOOSE 👇", "CHOSE 👇"}


_READ = object()


def item(ws, r, c, slot=None, text=_READ):
    if text is _READ:
        text = val(ws, r, c)
    return {"cell": f"{CL(CI(c) if isinstance(c, str) else c)}{r}", "row": r, "slot": slot,
            "text": None if text in PLACEHOLDER else text}


# ---------------------------------------------------------------- BYOP templates
BYOP_SPECS = {
    # tab: (day_col, purpose_col, category_col, exercise_col, [(day_label_cell, r0, r1, kind)])
    "AESTHETICS - FULL BODY - 4 DAYS": ("A", "B", "C", "D", [
        ("A3", 3, 9, "main"), ("A10", 10, 16, "main"), ("A17", 18, 22, "optional"),
        ("A24", 24, 30, "main"), ("A31", 31, 37, "main"), ("A38", 39, 43, "optional")]),
    "3 DAY BUILDER": ("A", "B", "C", "D", [
        ("A3", 3, 7, "main"), ("A8", 8, 12, "optional"), ("A13", 13, 17, "main"),
        ("A18", 18, 22, "optional"), ("A23", 23, 27, "main"), ("A28", 28, 32, "optional")]),
    "PUSHPULLLEGS - 3 or 6 DAYS A WK": ("B", "C", "D", "E", [
        ("B3", 3, 9, "main"), ("B10", 10, 16, "main"), ("B17", 17, 23, "main"),
        ("B24", 24, 30, "main"), ("B31", 31, 37, "main"), ("B38", 38, 44, "main")]),
    "ATHLETIC - FULL BODY - 4 DAYS A": ("B", "C", "D", "E", [
        ("B3", 3, 12, "main"), ("B13", 13, 22, "main"), ("B23", 24, 28, "optional"),
        ("B30", 30, 39, "main"), ("B40", 40, 49, "main"), ("B50", 51, 55, "optional")]),
    "JUSTICE LEAGUE": ("B", "C", "D", "E", [
        ("B3", 3, 12, "main"), ("B13", 13, 22, "main"), ("B23", 24, 28, "optional"),
        ("B30", 30, 39, "main"), ("B40", 40, 49, "main"), ("B50", 51, 55, "optional")]),
    "BOOTCAMP - OPTION 1": ("B", "C", "D", "E", [
        ("B3", 3, 9, "main"), ("B10", 10, 16, "main"), ("B17", 18, 22, "optional"),
        ("B24", 24, 30, "main"), ("B31", 31, 37, "main"), ("B38", 39, 43, "optional")]),
    "BOOTCAMP - OPTION 2": ("B", "C", "D", "E", [
        ("B3", 3, 7, "main"), ("B8", 8, 12, "main"), ("B13", 13, 17, "main"),
        ("B18", 18, 22, "main"), ("B23", 23, 27, "main"), ("B28", 28, 32, "main")]),
    "NO ACCESS - AESTHETICS - FB - 4": ("B", "C", "D", "E", [
        ("B3", 3, 9, "main"), ("B10", 10, 16, "main"), ("B17", 18, 22, "optional"),
        ("B24", 24, 30, "main"), ("B31", 31, 37, "main"), ("B38", 39, 43, "optional")]),
    "NO ACCESS - ATHLETIC - FB - 4 t": ("B", "C", "D", "E", [
        ("B3", 3, 12, "main"), ("B13", 13, 22, "main"), ("B23", 24, 28, "optional"),
        ("B30", 30, 39, "main"), ("B40", 40, 49, "main"), ("B50", 51, 55, "optional")]),
}
BYOP_TAB_NO = {
    "AESTHETICS - FULL BODY - 4 DAYS": "00", "3 DAY BUILDER": "01", "PUSHPULLLEGS - 3 or 6 DAYS A WK": "02",
    "ATHLETIC - FULL BODY - 4 DAYS A": "03", "JUSTICE LEAGUE": "04", "BOOTCAMP - OPTION 1": "05",
    "BOOTCAMP - OPTION 2": "06", "NO ACCESS - AESTHETICS - FB - 4": "07", "NO ACCESS - ATHLETIC - FB - 4 t": "08"}


def extract_byop():
    out = []
    w = wb("BYOP")
    for tab, (dc, pc, cc, ec, days) in BYOP_SPECS.items():
        ws = w[tab]
        for di, (dcell, r0, r1, kind) in enumerate(days):
            m = re.match(r"([A-Z]+)(\d+)", dcell)
            dlabel = val(ws, int(m.group(2)), m.group(1))
            if dlabel is None:
                dlabel = "(blank day label; 4th main day, between WEDNESDAY and FRIDAY)"
            items = []
            for r in range(r0, r1 + 1):
                purpose, cat = val(ws, r, pc), val(ws, r, cc)
                if purpose is None and cat is None:
                    continue
                it = item(ws, r, ec, slot={"PURPOSE": purpose, "CATEGORY": cat})
                it["purpose_cell"], it["category_cell"] = f"{pc}{r}", f"{cc}{r}"
                items.append(it)
            out.append({"sid": f"BYOP/{BYOP_TAB_NO[tab]}/{dcell}", "book": "BYOP", "tab": tab,
                        "tab_no": BYOP_TAB_NO[tab], "track": None, "week": "WEEK 1 (template)",
                        "day": dlabel, "day_cell": dcell, "day_index": di, "level": None,
                        "kind": "template_" + kind, "items": items})
    return out


# ---------------------------------------------------------------- MASTER MODEL
def col_session(ws, sid, tab, tab_no, col, r0, r1, day, week=None, level=None, track=None, kind="filled",
                skip=()):
    items = []
    for r in range(r0, r1 + 1):
        t = val(ws, r, col)
        if t is None or t in skip:
            continue
        items.append(item(ws, r, col, text=t))
    return {"sid": sid, "book": "MASTER_MODEL", "tab": tab, "tab_no": tab_no, "track": track, "week": week,
            "day": day, "level": level, "kind": kind, "items": items}


def extract_mm00():
    tab = "EXAMPLE UPPERLOWER SPLIT"; ws = wb("MASTER_MODEL")[tab]; out = []
    for col, lev in (("C", "B10 'NOVICE'"), ("G", "F10 'NEWBY'"), ("K", "J10 'INTERMEDIATE'")):
        out.append(col_session(ws, f"MM/00/{col}13", tab, "00", col, 13, 23, "A11 'LOWER BODY WORKOUT' (ADAY)", level=lev))
        out.append(col_session(ws, f"MM/00/{col}27", tab, "00", col, 27, 35, "A25 'UPPER BODY WORKOUT' (BDAY)", level=lev))
    return out


def extract_mm06():
    tab = "EXAMPLE PUSHPULLLEGS SPLIT"; ws = wb("MASTER_MODEL")[tab]; out = []
    for col, lev in (("C", "A10 'NOVICE'"), ("G", "E10 'NEWBY'"), ("K", "I10 'INTERMEDIATE'")):
        out.append(col_session(ws, f"MM/06/{col}14", tab, "06", col, 14, 20, "A12 'PUSH DAY WORKOUT' (ADAY)", level=lev))
        out.append(col_session(ws, f"MM/06/{col}25", tab, "06", col, 25, 31, "A23 'PULL DAY WORKOUT' (BDAY)", level=lev))
        out.append(col_session(ws, f"MM/06/{col}36", tab, "06", col, 36, 46, "A34 'LEG DAY WORKOUT' (CDAY)", level=lev))
    return out


def extract_mm01():
    tab = "EXAMPLE FAT BURNBODY TONING PRO"; ws = wb("MASTER_MODEL")[tab]; out = []
    wk = {"D": "WK 1", "G": "WK 2", "J": "WK 3", "M": "WK 4"}
    for col in "DGJM":
        for (r0, r1, day) in ((7, 11, "A8 'MONDAY'"), (12, 16, "A14 'TUESDAY'"), (17, 17, "A18 'WEDNESDAY'"),
                              (20, 24, "A22 'THURSDAY'"), (25, 29, "A27 'FRIDAY'"), (30, 30, "A31 'SATURDAY'")):
            s = col_session(ws, f"MM/01/{col}{r0}", tab, "01", col, r0, r1, day, week=wk[col],
                            track="B4 'MUSCLE BUILDING'")
            if r0 in (17, 30):
                s["kind"] = "cardio_day"
            out.append(s)
    # circuits ('CIRCUITS RECCOMENDED'): stations listed in the week column that is filled
    for (r0, r1, day, cols) in ((39, 42, "A40 'MONDAY'", "DJ"), (44, 47, "A46 'TUESDAY'", "GM"),
                                (52, 55, "A54 'THURSDAY'", "GM"), (57, 60, "A59 'FRIDAY'", "DJ")):
        for col in cols:
            s = col_session(ws, f"MM/01/{col}{r0}", tab, "01", col, r0, r1, day, week=wk[col],
                            track="B36 'CIRCUITS RECCOMENDED'", kind="circuit")
            out.append(s)
    return out


def extract_mm05():
    tab = "EXAMPLE ATHLETIC FULL BODY SPLI"; ws = wb("MASTER_MODEL")[tab]; out = []
    specs = [("C", "A9 'STARTING OFF SEASON WORKOUTS'", [(12, 21), (24, 30), (33, 39)]),
             ("F", "D9 ' MID OFF SEASON WORKOUTS'", [(12, 21), (24, 37), (41, 48), (51, 67)]),
             ("J", "H9 ' ADVANCED OFF SEASON WORKOUTS'", [(12, 17), (20, 25), (28, 35), (38, 43)])]
    for col, lev, blocks in specs:
        for i, (r0, r1) in enumerate(blocks, 1):
            out.append(col_session(ws, f"MM/05/{col}{r0}", tab, "05", col, r0, r1, f"WORKOUT {i}", level=lev,
                                   skip=("WORKOUT", "SETS", "REPS", "EXERCISE")))
    return out


def extract_mm07():
    tab = "EXAMPLE BOXING 8 WEEK BUILD UP"; ws = wb("MASTER_MODEL")[tab]; out = []
    days = [("E4 'MONDAY OFF SEASON WORK'", 7, 14, 7, 14), ("E16 'TUESDAY OFF SEASON WORK'", 19, 25, 19, 23),
            ("E28 'THURSDAY OFF SEASON WORK'", 31, 37, 31, 37), ("E41 'FRIDAY OFF SEASON WORK'", 44, 49, 44, 48)]
    for day, a0, a1, b0, b1 in days:
        out.append(col_session(ws, f"MM/07/D{a0}", tab, "07", "D", a0, a1, day, week="Week 1-2 (col D)"))
        out.append(col_session(ws, f"MM/07/L{b0}", tab, "07", "L", b0, b1, day, week="Week 3-8 (col L)"))
    return out


def extract_mm12():
    tab = "(NOT DONE) EXAMPLE CONJUGATE"; ws = wb("MASTER_MODEL")[tab]; out = []
    for col in "EFGHIJK":
        wk = val(ws, 12, col)
        for (r0, r1, day) in ((13, 16, "A13 'MONDAY'"), (17, 20, "A17 'WEDNESDAY'"), (21, 24, "A21 'THURSDAY'"),
                              (25, 27, "A25 'SATURDAY'")):
            s = col_session(ws, f"MM/12/{col}{r0}", tab, "12", col, r0, r1, day, week=wk)
            for it in s["items"]:
                it["slot"] = {"M": val(ws, it["row"], "M"), "LOAD": val(ws, it["row"], "D")}
            out.append(s)
    return out


def extract_mm04():
    tab = "(NOT DONE) 4 TRACKS PROGRAMMING"; ws = wb("MASTER_MODEL")[tab]; out = []
    # STRENGTH track, left (8-week, E-K; L = TEST/DELOAD) and right (12-week, O-Z)
    sdays = ((9, 12, "A9 'MONDAY'"), (13, 16, "A13 'WEDNESDAY'"), (17, 20, "A17 'THURSDAY'"), (21, 24, "A21 'SATURDAY'"))
    for col in "EFGHIJK":
        for r0, r1, day in sdays:
            s = col_session(ws, f"MM/04/STR-L/{col}{r0}", tab, "04", col, r0, r1, day, week=val(ws, 8, col),
                            track="A3 'STRENGTH' left (E-L, 'conjugate style training')")
            for it in s["items"]:
                it["slot"] = {"M": val(ws, it["row"], "M"), "LOAD": val(ws, it["row"], "D")}
            out.append(s)
    for col in [CL(i) for i in range(CI("O"), CI("Z") + 1)]:
        for r0, r1, day in sdays:
            s = col_session(ws, f"MM/04/STR-R/{col}{r0}", tab, "04", col, r0, r1, day, week=val(ws, 8, col),
                            track="A3 'STRENGTH' right (O-Z, 12 wk)")
            for it in s["items"]:
                it["slot"] = {"M": val(ws, it["row"], "M"), "N": val(ws, it["row"], "N")}
            out.append(s)
    # AESTHETICS track left (E-K) - upper/lower x2 + cardio Wed/Sat
    adays = ((35, 40, "A35 'MONDAY'"), (41, 45, "A41 'TUESDAY'"), (46, 46, "A46 'WEDNESDAY '"),
             (49, 54, "A49 'THURSDAY'"), (55, 59, "A55 'FRIDAY'"), (60, 60, "A60 'SATURDAY'"))
    for col in "EFGHIJK":
        for r0, r1, day in adays:
            s = col_session(ws, f"MM/04/AES-L/{col}{r0}", tab, "04", col, r0, r1, day, week=val(ws, 34, col),
                            track="A30 'AESTHETICS' left (E-L, 'bodybuilding style')")
            if r0 in (46, 60):
                s["kind"] = "cardio_day"
            out.append(s)
    # AESTHETICS track right (O-Z) - PPL x2
    rdays = ((35, 40, "N35 'MONDAY' (M38 'PUSH')"), (41, 45, "N41 'TUESDAY' (M43 'LEGS')"),
             (46, 50, "N46 'WEDNESDAY ' (M49 'PULL')"), (51, 56, "N51 'THURSDAY' (M54 'PUSH')"),
             (57, 61, "N57 'FRIDAY' (legs)"), (62, 66, "N62 'SATURDAY' (M64 'PULL')"))
    for col in [CL(i) for i in range(CI("O"), CI("Z") + 1)]:
        for r0, r1, day in rdays:
            s = col_session(ws, f"MM/04/AES-R/{col}{r0}", tab, "04", col, r0, r1, day,
                            week=f"{val(ws, 34, col)} ({val(ws, 33, col) or ''})".replace(" ()", ""),
                            track="A30 'AESTHETICS' right (O-Z, 12 wk PPL)")
            out.append(s)
    # HYBRID track (E-L) - slot rows with B/C/D labels
    hdays = ((74, 79, "A74 'MONDAY'"), (80, 83, "A80 'TUESDAY'"), (86, 90, "A86 'THURSDAY'"),
             (92, 96, "A92 'FRIDAY'"), (98, 101, "A98 'SATURDAY'"))
    for col in "EFGHIJKL":
        for r0, r1, day in hdays:
            items = []
            for r in range(r0, r1 + 1):
                b, c, d, x = val(ws, r, "B"), val(ws, r, "C"), val(ws, r, "D"), val(ws, r, col)
                if isinstance(x, (int, float)):
                    x = None  # rows 78-79 hold numbers (sets) not exercises
                if b is None and d is None and x is None:
                    continue
                if b is None and c is None and d is None:
                    continue
                it = item(ws, r, col, slot={"B": b, "C": c, "D": d}, text=x)
                if c is None and d is None and x is None:
                    it["empty_template_row"] = True  # e.g. B99 'OLY/ POWER' on SATURDAY: label only, no content
                items.append(it)
            out.append({"sid": f"MM/04/HYB/{col}{r0}", "book": "MASTER_MODEL", "tab": tab, "tab_no": "04",
                        "track": "A69 'HYBRID' ('strenght - athletic - crossfit')", "week": val(ws, 73, col),
                        "day": day, "level": None, "kind": "slot_grid", "items": items})
    return out


def extract_mm02():
    tab = "BASEBALL"; ws = wb("MASTER_MODEL")[tab]; out = []
    blocks = [("A24 'BLOCK 1'", 26, 27, 40, [22, 25]), ("A41 'BLOCK 2'", 42, 45, 58, [42, 43]),
              ("A59 'BLOCK 3'", 61, 62, 76, [60, 61])]
    days = ("MONDAY", "TUESDAY", "THURSDAY", "FRIDAY")
    for bname, hr0, hr1, rend, wkrows in blocks:
        # week labels in this block
        wlabels = []
        for r in wkrows:
            for c in range(1, ws.max_column + 1):
                v = val(ws, r, c)
                if isinstance(v, str) and re.match(r"(?i)^(wk|week|deload)", v):
                    wlabels.append((c, f"{CL(c)}{r} '{v}'"))
        wlabels.sort()
        heads = []
        for r in range(hr0, hr1 + 1):
            for c in range(1, ws.max_column + 1):
                v = val(ws, r, c)
                if isinstance(v, str) and v.upper() in days:
                    heads.append((r, c, v))
        for (hr, hc, dname) in heads:
            # locate the 'Sets' header in the row below to find the column layout
            setc = None
            for c in range(hc - 1, hc + 4):
                if val(ws, hr + 1, c) == "Sets":
                    setc = c
                    break
            if setc is None:
                raise RuntimeError(f"no Sets header near {CL(hc)}{hr}")
            labc, exc = setc - 1, setc + 2
            wk = [lab for (c, lab) in wlabels if c <= hc + 1]
            wk = wk[-1] if wk else "?"
            items = []
            for r in range(hr + 2, rend + 1):
                lab, ex = val(ws, r, labc), val(ws, r, exc)
                if ex is None:
                    continue
                if isinstance(ex, (int, float)):
                    continue
                items.append({"cell": f"{CL(exc)}{r}", "row": r, "slot": lab, "text": ex,
                              "label_cell": f"{CL(labc)}{r}"})
            out.append({"sid": f"MM/02/{CL(hc)}{hr}", "book": "MASTER_MODEL", "tab": tab, "tab_no": "02",
                        "track": bname, "week": wk, "day": f"{CL(hc)}{hr} '{dname}'", "level": None,
                        "kind": "filled", "items": items, "label_col": CL(labc), "ex_col": CL(exc)})
    return out


def extract_mm11():
    """CROSSFIT: no exercise lists; returns category-level weekly grids (for R11 and category order)."""
    tab = "EXAMPLE CROSSFIT SPLIT"; ws = wb("MASTER_MODEL")[tab]
    grids = []
    for r in range(27, 146):
        month = val(ws, r, "A")
        if month and val(ws, r + 1, "B") == "MON":
            g = {"month_cell": f"A{r}", "month": month, "rows": {}}
            for rr in range(r + 2, r + 9):
                cat = val(ws, rr, "A")
                g["rows"][cat] = {val(ws, r + 1, c): val(ws, rr, c) for c in "BCDEFGH"}
            grids.append(g)
    class_model = {"R3": val(ws, 3, "R"), "S3": val(ws, 3, "S"), "T3": val(ws, 3, "T")}
    class_split = {}
    for rr in range(13, 17):
        class_split[val(ws, rr, "P")] = {val(ws, 12, c): val(ws, rr, c) for c in "QRSTUV"}
    return {"monthly_grids": grids, "class_model_header": class_model, "class_split": class_split}


def extract_all():
    sessions = []
    sessions += extract_byop()
    sessions += extract_mm00() + extract_mm01() + extract_mm02() + extract_mm04()
    sessions += extract_mm05() + extract_mm06() + extract_mm07() + extract_mm12()
    return sessions


if __name__ == "__main__":
    ss = extract_all()
    from collections import Counter
    print(len(ss), Counter((s["book"], s["tab"]) for s in ss))
