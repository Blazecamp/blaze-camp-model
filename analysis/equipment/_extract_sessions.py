# -*- coding: utf-8 -*-
"""Extract every session / line item from the MASTER MODEL program tabs (read-only).

Source: raw/MASTER_MODEL.xlsx (openpyxl, values only). Nothing in Google Sheets is touched.
Output: a list of session dicts; each item keeps its source cell so every claim can be cited.
"""
import datetime, re
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

RAW = 'raw/MASTER_MODEL.xlsx'

def fmt(v):
    """Display value the way Google Sheets shows it. m/d dates were typed as 'N/N' reps-per-side."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return f"{v.month}/{v.day}"
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    if isinstance(v, int):
        return str(v)
    s = str(v)
    return s if s.strip() != '' else None

class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and fmt(c.value) is not None:
                    self.cells[(c.row, c.column)] = c.value
    def v(self, r, c):
        if isinstance(c, str):
            c = CI(c)
        return fmt(self.cells.get((r, c)))
    def raw(self, r, c):
        if isinstance(c, str):
            c = CI(c)
        return self.cells.get((r, c))

def cell(r, c):
    if isinstance(c, int):
        c = L(c)
    return f"{c}{r}"

def item(sh, r, cs, cr, ce, part=None, slot=None, extra=None):
    name = sh.v(r, ce)
    if name is None:
        return None
    d = {
        'cell': cell(r, ce),
        'row': r,
        'sets_raw': sh.v(r, cs) if cs else None,
        'reps_raw': sh.v(r, cr) if cr else None,
        'name': name.strip() if isinstance(name, str) else name,
        'part': part,
        'slot': slot,
    }
    if extra:
        d.update(extra)
    return d

def mk_session(program, tab, population, variant, week, day, items, notes=None, source_range=None, **kw):
    s = {'program': program, 'tab': tab, 'population': population, 'variant': variant,
         'week': week, 'day': day, 'items': [i for i in items if i], 'notes': notes or [],
         'source_range': source_range}
    s.update(kw)
    for k, it in enumerate(s['items'], 1):
        it['pos'] = k
    return s

# ---------------------------------------------------------------- 00 UPPER/LOWER
def ex_ul(wb):
    sh = Sheet(wb['EXAMPLE UPPERLOWER SPLIT'])
    out = []
    levels = [('NOVICE', 'A', 'B', 'C'), ('NEWBY', 'E', 'F', 'G'), ('INTERMEDIATE', 'I', 'J', 'K')]
    sess = [('A', 'LOWER BODY WORKOUT (ADAY)', 13, 23), ('B', 'UPPER BODY WORKOUT (BDAY)', 27, 35)]
    for lvl, cs, cr, ce in levels:
        for code, nm, r0, r1 in sess:
            its = [item(sh, r, CI(cs), CI(cr), CI(ce)) for r in range(r0, r1 + 1)]
            out.append(mk_session('UPPER/LOWER', sh.title, 'gen_pop', lvl, 'all weeks (no week columns)',
                                  nm, its, source_range=f"{cs}{r0}:{ce}{r1}", day_code=code))
    return out

# ---------------------------------------------------------------- 06 PPL
def ex_ppl(wb):
    sh = Sheet(wb['EXAMPLE PUSHPULLLEGS SPLIT'])
    out = []
    levels = [('NOVICE', 'A', 'B', 'C'), ('NEWBY', 'E', 'F', 'G'), ('INTERMEDIATE', 'I', 'J', 'K')]
    sess = [('A', 'PUSH DAY WORKOUT (ADAY)', 14, 20), ('B', 'PULL DAY WORKOUT (BDAY)', 25, 31),
            ('C', 'LEG DAY WORKOUT (CDAY)', 36, 46)]
    for lvl, cs, cr, ce in levels:
        for code, nm, r0, r1 in sess:
            its = [item(sh, r, CI(cs), CI(cr), CI(ce)) for r in range(r0, r1 + 1)]
            out.append(mk_session('PUSH/PULL/LEGS', sh.title, 'gen_pop', lvl, 'all weeks (no week columns)',
                                  nm, its, source_range=f"{cs}{r0}:{ce}{r1}", day_code=code))
    return out

# ---------------------------------------------------------------- 01 FAT BURN
def ex_fatburn(wb):
    sh = Sheet(wb['EXAMPLE FAT BURNBODY TONING PRO'])
    out = []
    weeks = [('WK 1', 'B', 'C', 'D'), ('WK 2', 'E', 'F', 'G'), ('WK 3', 'H', 'I', 'J'), ('WK 4', 'K', 'L', 'M')]
    # MUSCLE BUILDING block (B4). Day labels sit in col A (A8 MONDAY, A14 TUESDAY, A18 WEDNESDAY,
    # A22 THURSDAY, A27 FRIDAY, A31 SATURDAY); rows 7-11 are all lower-body, 12-16 all upper-body.
    days = [('MONDAY', 7, 11), ('TUESDAY', 12, 16), ('WEDNESDAY', 17, 17), ('THURSDAY', 20, 24),
            ('FRIDAY', 25, 29), ('SATURDAY', 30, 30)]
    for wk, cs, cr, ce in weeks:
        for day, r0, r1 in days:
            its = [item(sh, r, CI(cs), CI(cr), CI(ce)) for r in range(r0, r1 + 1)]
            out.append(mk_session('FAT BURN / BODY TONING', sh.title, 'gen_pop', 'MUSCLE BUILDING', wk, day, its,
                                  source_range=f"{cs}{r0}:{ce}{r1}",
                                  notes=['day-to-row mapping INFERRED from col-A labels and lower/upper content']))
    # CIRCUITS RECCOMENDED block (B36). Only the first station carries SETS -> rounds for the circuit.
    cdays = [('MONDAY', 39, 42), ('TUESDAY', 44, 47), ('THURSDAY', 52, 55), ('FRIDAY', 57, 60)]
    for wk, cs, cr, ce in weeks:
        for day, r0, r1 in cdays:
            its = [item(sh, r, CI(cs), CI(cr), CI(ce)) for r in range(r0, r1 + 1)]
            its = [i for i in its if i]
            if not its:
                continue
            rounds = its[0]['sets_raw']
            for i in its[1:]:
                i['paired_with_prev'] = True
            out.append(mk_session('FAT BURN / BODY TONING', sh.title, 'gen_pop', 'CIRCUITS RECCOMENDED', wk, day, its,
                                  source_range=f"{cs}{r0}:{ce}{r1}", circuit_rounds=rounds,
                                  notes=['timed 4-station circuit; only first station has SETS (rounds)']))
    return out

# ---------------------------------------------------------------- 05 ATHLETIC FULL BODY
def ex_athletic(wb):
    sh = Sheet(wb['EXAMPLE ATHLETIC FULL BODY SPLI'])
    out = []
    groups = [('STARTING OFF SEASON WORKOUTS', 'A', 'B', 'C', 12, 39),
              (' MID OFF SEASON WORKOUTS', 'D', 'E', 'F', 12, 67),
              (' ADVANCED OFF SEASON WORKOUTS', 'H', 'I', 'J', 12, 43)]
    for gname, cs, cr, ce, r0, r1 in groups:
        cur, wk_no, start = [], 1, r0
        section = None
        def flush(end):
            nonlocal cur, wk_no
            if cur:
                out.append(mk_session('ATHLETIC FULL BODY', sh.title, 'athlete', gname.strip(), 'off-season (no week columns)',
                                      f'WORKOUT {wk_no}', cur, source_range=f"{cs}{start}:{ce}{end}"))
            cur = []
            wk_no += 1
        for r in range(r0, r1 + 1):
            s, e = sh.v(r, cs), sh.v(r, ce)
            if s == 'WORKOUT':
                flush(r - 1)
                start = r + 2
                section = None
                continue
            if s in ('SETS',):
                continue
            if s and s.strip('* ') == 'MEDBALLS':      # D24 / D51 section label '**MEDBALLS**'
                section = 'MEDBALLS'
                continue
            it = item(sh, r, CI(cs), CI(cr), CI(ce), part=section)
            if it:
                cur.append(it)
        flush(r1)
    return out

# ---------------------------------------------------------------- 07 BOXING
def ex_boxing(wb):
    sh = Sheet(wb['EXAMPLE BOXING 8 WEEK BUILD UP'])
    out = []
    days = [('MONDAY OFF SEASON WORK', 7, 14), ('TUESDAY OFF SEASON WORK', 19, 25),
            ('THURSDAY OFF SEASON WORK', 31, 37), ('FRIDAY OFF SEASON WORK', 44, 49)]
    for day, r0, r1 in days:
        # phase 1: weeks 1-2 -> sets A/B, reps C, exercise D
        its = [item(sh, r, CI('A'), CI('C'), CI('D'), extra={'sets_by_week': {'Week 1': sh.v(r, 'A'), 'Week 2': sh.v(r, 'B')}})
               for r in range(r0, r1 + 1)]
        out.append(mk_session('BOXING 8 WEEK BUILD UP', sh.title, 'boxing', 'Weeks 1-2', 'Week 1-2', day, its,
                              source_range=f"A{r0}:D{r1}"))
        its = [item(sh, r, CI('E'), CI('K'), CI('L'), extra={'sets_by_week': {f'Week {w}': sh.v(r, c) for w, c in zip(range(3, 9), 'EFGHIJ')}})
               for r in range(r0, r1 + 1)]
        out.append(mk_session('BOXING 8 WEEK BUILD UP', sh.title, 'boxing', 'Weeks 3-8', 'Week 3-8', day, its,
                              source_range=f"E{r0}:L{r1}"))
    return out

def boxing_medball_list(wb):
    sh = Sheet(wb['EXAMPLE BOXING 8 WEEK BUILD UP'])
    return [{'cell': f'G{r}', 'reps': sh.v(r, 'F'), 'name': sh.v(r, 'G')} for r in range(51, 70) if sh.v(r, 'G')]

# ---------------------------------------------------------------- 12 CONJUGATE
def ex_conjugate(wb):
    sh = Sheet(wb['(NOT DONE) EXAMPLE CONJUGATE'])
    out = []
    pink = {f'WK{i+1}': sh.v(9, c) for i, c in enumerate('EFGHIJKL')}   # EAD1DC reps row (max effort)
    blue = {f'WK{i+1}': sh.v(10, c) for i, c in enumerate('EFGHIJKL')}  # CFE2F3 reps row (dynamic effort)
    days = [('MONDAY', 13, 16), ('WEDNESDAY', 17, 20), ('THURSDAY', 21, 24), ('SATURDAY', 25, 28)]
    for wi, col in enumerate('EFGHIJKL'):
        wk = f'WK{wi+1}'
        for day, r0, r1 in days:
            its = []
            for r in range(r0, r1 + 1):
                reps = sh.v(r, 'C')
                if reps == 'SEE':
                    fill = sh.ws[f'C{r}'].fill.fgColor.rgb
                    reps_res = (pink if str(fill).endswith('EAD1DC') else blue)[wk]
                    reps = f"SEE -> {reps_res} (row {'9' if str(fill).endswith('EAD1DC') else '10'})"
                it = item(sh, r, CI('B'), None, CI(col), extra={'reps_raw': reps, 'load_raw': sh.v(r, 'D'),
                                                               'role_label': sh.v(r, 'M')})
                if it:
                    its.append(it)
            out.append(mk_session('CONJUGATE (NOT DONE)', sh.title, 'strength_powerlifting', '8 WEEK CYCLE', wk, day, its,
                                  source_range=f"B{r0}:M{r1} col {col}"))
    return out

# ---------------------------------------------------------------- 04 FOUR TRACKS
def ex_tracks(wb):
    sh = Sheet(wb['(NOT DONE) 4 TRACKS PROGRAMMING'])
    out = []
    # ---- STRENGTH track, 8-week conjugate columns E..L
    green = {f'WK{i+1}': sh.v(4, c) for i, c in enumerate('EFGHIJKL')}
    pink = {f'WK{i+1}': sh.v(5, c) for i, c in enumerate('EFGHIJKL')}
    blue = {f'WK{i+1}': sh.v(6, c) for i, c in enumerate('EFGHIJKL')}
    days = [('MONDAY', 9, 12), ('WEDNESDAY', 13, 16), ('THURSDAY', 17, 20), ('SATURDAY', 21, 24)]
    for wi, col in enumerate('EFGHIJKL'):
        wk = f'WK{wi+1}'
        for day, r0, r1 in days:
            its = []
            for r in range(r0, r1 + 1):
                sets = sh.v(r, 'B')
                if sets and str(sh.ws[f'B{r}'].fill.fgColor.rgb).endswith('B6D7A8'):
                    sets = f"{sets} (green; row-4 value {green[wk]})"
                reps = sh.v(r, 'C')
                if reps == 'SEE':
                    fill = str(sh.ws[f'C{r}'].fill.fgColor.rgb)
                    reps = f"SEE -> {(pink if fill.endswith('EAD1DC') else blue)[wk]} (row {'5' if fill.endswith('EAD1DC') else '6'})"
                it = item(sh, r, None, None, CI(col), extra={'sets_raw': sets, 'reps_raw': reps, 'load_raw': sh.v(r, 'D'),
                                                             'role_label': sh.v(r, 'M')})
                if it:
                    its.append(it)
            out.append(mk_session('4 TRACKS - STRENGTH (conjugate style)', sh.title, 'strength_powerlifting',
                                  '8-week columns E:L', wk, day, its, source_range=f"B{r0}:M{r1} col {col}"))
    # ---- STRENGTH track, 12-week columns O..Z (N = prescription label, rows 3-6 = %/RPE/SETS/REPS)
    for wi, col in enumerate('OPQRSTUVWXYZ'):
        wk = sh.v(8, col)
        hdr = {'%': sh.v(3, col), 'RPE': sh.v(4, col), 'SETS': sh.v(5, col), 'REPS': sh.v(6, col)}
        for day, r0, r1 in days:
            its = []
            for r in range(r0, r1 + 1):
                it = item(sh, r, None, None, CI(col), extra={'reps_raw': sh.v(r, 'N'), 'role_label': sh.v(r, 'M'),
                                                             'week_header': hdr})
                if it:
                    its.append(it)
            out.append(mk_session('4 TRACKS - STRENGTH (conjugate style)', sh.title, 'strength_powerlifting',
                                  '12-week columns O:Z', wk, day, its, source_range=f"N{r0}:N{r1} + col {col}"))
    # ---- AESTHETICS track, 8-week columns E..L (sets row 32, reps row 33)
    sets_w = {c: sh.v(32, c) for c in 'EFGHIJK'}
    reps_w = {c: sh.v(33, c) for c in 'EFGHIJK'}
    adays = [('MONDAY', 35, 40), ('TUESDAY', 41, 45), ('WEDNESDAY', 46, 46), ('THURSDAY', 49, 54),
             ('FRIDAY', 55, 59), ('SATURDAY', 60, 60)]
    for wi, col in enumerate('EFGHIJKL'):
        wk = sh.v(34, col)
        for day, r0, r1 in adays:
            its = []
            for r in range(r0, r1 + 1):
                it = item(sh, r, None, None, CI(col), extra={'sets_raw': (sh.v(r, 'B') or sets_w.get(col)),
                                                             'reps_raw': reps_w.get(col), 'load_raw': sh.v(r, 'D')})
                if it:
                    its.append(it)
            if its:
                out.append(mk_session('4 TRACKS - AESTHETICS (bodybuilding style)', sh.title, 'gen_pop',
                                      '8-week columns E:L', wk, day, its, source_range=f"{col}{r0}:{col}{r1}"))
    # ---- AESTHETICS track, 12-week columns O..Z (N = day labels, M = PUSH/LEGS/PULL)
    a12 = [('MONDAY', 35, 40), ('TUESDAY', 41, 45), ('WEDNESDAY', 46, 50), ('THURSDAY', 51, 56),
           ('FRIDAY', 57, 61), ('SATURDAY', 62, 66)]
    mlabel = {38: 'PUSH', 43: 'LEGS', 49: 'PULL', 54: 'PUSH', 64: 'PULL'}
    for col in 'OPQRSTUVWXYZ':
        wk = sh.v(34, col)
        hdr = {'SETS': sh.v(31, col), 'REPS': sh.v(32, col), 'A/B': sh.v(33, col)}
        for day, r0, r1 in a12:
            its = [item(sh, r, None, None, CI(col), extra={'sets_raw': hdr['SETS'], 'reps_raw': hdr['REPS'], 'week_header': hdr})
                   for r in range(r0, r1 + 1)]
            lab = [mlabel[r] for r in range(r0, r1 + 1) if r in mlabel]
            out.append(mk_session('4 TRACKS - AESTHETICS (bodybuilding style)', sh.title, 'gen_pop',
                                  '12-week columns O:Z', wk, day, its, source_range=f"{col}{r0}:{col}{r1}",
                                  split_label=(lab[0] if lab else None)))
    # ---- HYBRID track rows 74-101, weeks E..L (sets row 71, reps row 72, AMRAP row 70)
    hdays = [('MONDAY', 74, 79), ('TUESDAY', 80, 83), ('THURSDAY', 86, 90), ('FRIDAY', 92, 96), ('SATURDAY', 98, 101)]
    for col in 'EFGHIJKL':
        wk = sh.v(73, col)
        hdr = {'AMRAP': sh.v(70, col), 'SETS': sh.v(71, col), 'REPS': sh.v(72, col)}
        for day, r0, r1 in hdays:
            its = []
            for r in range(r0, r1 + 1):
                nm = sh.v(r, col)
                tmpl = sh.v(r, 'D')
                if nm is None or re.fullmatch(r'[0-9.]+', nm or ''):
                    nm2 = tmpl
                    if nm2 is None:
                        continue
                    it = {'cell': f"D{r}", 'row': r, 'name': nm2, 'sets_raw': nm if nm else None,
                          'reps_raw': None, 'part': sh.v(r, 'B'), 'slot': sh.v(r, 'C'), 'template_only': True}
                else:
                    is_str = (sh.v(r, 'B') or '').strip().upper() == 'STRENGTH'
                    it = {'cell': f"{col}{r}", 'row': r, 'name': nm,
                          'sets_raw': hdr['SETS'] if is_str else None, 'reps_raw': hdr['REPS'] if is_str else None,
                          'part': sh.v(r, 'B'), 'slot': sh.v(r, 'C'), 'template_name': tmpl,
                          'header_note': 'rows 70-72 (AMRAP/SETS/REPS) attached to STRENGTH rows only (INFERRED)'}
                its.append(it)
            out.append(mk_session('4 TRACKS - HYBRID (strength-athletic-crossfitish)', sh.title, 'athlete_hybrid',
                                  'weeks E:L', wk, day, its, source_range=f"B{r0}:{col}{r1}", week_header=hdr))
    return out

# ---------------------------------------------------------------- 11 CROSSFIT (component grid, no exercises)
def ex_crossfit(wb):
    sh = Sheet(wb['EXAMPLE CROSSFIT SPLIT'])
    out = []
    comps = ['AEROBIC C.', 'SKILL', 'STRENGTH', 'POWER', 'RESISTANCE', 'WODs', 'TEST/PEAK']
    r = 27
    while r <= 146:
        m = sh.v(r, 'A')
        if m and sh.v(r + 1, 'B') == 'MON':
            for ci, col in enumerate('BCDEFGH'):
                day = sh.v(r + 1, col)
                its = []
                for k, comp in enumerate(comps):
                    rr = r + 2 + k
                    val = sh.v(rr, col)
                    if val and val != 'OFF':
                        its.append({'cell': f'{col}{rr}', 'row': rr, 'name': f"{comp}: {val}", 'component': comp,
                                    'component_value': val, 'sets_raw': None, 'reps_raw': None,
                                    'level_I': sh.v(rr, 'I')})
                offs = [sh.v(r + 2 + k, col) for k in range(len(comps))]
                if its:
                    out.append(mk_session('CROSSFIT SPLIT', sh.title, 'crossfit', 'COMPETE CROSSFIT SPLIT (monthly grid)',
                                          m.strip(), day, its, source_range=f"{col}{r+2}:{col}{r+8}"))
            r += 9
        else:
            r += 1
    # class model P12:V16
    for col in 'QRSTUV':
        day = sh.v(12, col)
        its = []
        # ordered by the class header R3:T3 'WARMUP/SKILL' | 'STRENGTH/POWER' | 'WOD' (aerobic capacity placed last, INFERRED)
        for rr, comp in [(15, 'SKILLS'), (13, 'STRNGTH'), (14, 'PWR'), (16, 'AEROBIC CAPACITY')]:
            val = sh.v(rr, col)
            if val:
                its.append({'cell': f'{col}{rr}', 'row': rr, 'name': f'{comp}: {val}', 'component': comp,
                            'component_value': val, 'sets_raw': None, 'reps_raw': None})
        out.append(mk_session('CROSSFIT SPLIT', sh.title, 'crossfit', 'CLASS PROGRAMMING SPLIT (P1)', 'template', day, its,
                              source_range=f'{col}12:{col}16',
                              notes=["class order header R3:T3 = 'WARMUP/SKILL' | 'STRENGTH/POWER' | 'WOD'"]))
    return out

# ---------------------------------------------------------------- 02 BASEBALL (generic grid parser)
DAYN = {'MONDAY', 'TUESDAY', 'THURSDAY', 'FRIDAY'}
PART_RE = re.compile(r'^\s*part\s+([a-e])\b', re.I)

def ex_baseball(wb):
    sh = Sheet(wb['BASEBALL'])
    blocks = [('BLOCK 1', 24, 40), ('BLOCK 2', 41, 58), ('BLOCK 3', 59, 76)]
    out = []
    for bname, b0, b1 in blocks:
        headers = sorted([(r, c) for (r, c), v in sh.cells.items()
                          if b0 <= r <= b1 and isinstance(v, str) and v.strip().lower() == 'sets'])
        # data columns used by any session in this block
        datacols = set()
        for r, c in headers:
            datacols |= {c, c + 1, c + 2}
        # part columns: any column with 'Part A' in the block
        partcols = sorted({c for (r, c), v in sh.cells.items() if b0 <= r <= b1 and isinstance(v, str)
                           and re.match(r'^\s*part\s+a\b', v, re.I) and c not in datacols})
        for hr, hc in headers:
            # day name: nearest day header above, within cols hc-1..hc+2
            day = None
            for dr in (1, 2):
                for dc in (0, -1, 1, 2, -2):
                    v = sh.v(hr - dr, hc + dc)
                    if v and v.strip().upper() in DAYN:
                        day = v.strip().upper(); break
                if day:
                    break
            # stop row: next 'Sets' header in same column, or block end
            nxt = [r for r, c in headers if c == hc and r > hr]
            stop = (min(nxt) - 2) if nxt else b1
            pc = max([p for p in partcols if p <= hc], default=None)
            slotc = hc - 1 if (hc - 1) not in datacols and (hc - 1) != pc else None
            # week label: search rows b0-4 .. hr for a week-ish label in cols pc..hc+2
            wk = None
            for rr in range(hr, b0 - 5, -1):
                for cc in range((pc or hc) - 1, hc + 3):
                    v = sh.v(rr, cc)
                    if v and re.match(r'^\s*(wk|week)\s*\d|^\s*DELOAD WEEK', v, re.I):
                        wk = v.strip(); break
                if wk:
                    break
            items, cur_part, own_part = [], None, None
            own_labels = {}
            for r in range(hr + 1, stop + 1):
                for cc in [hc] + ([slotc] if slotc else []):
                    v = sh.v(r, cc)
                    if v and PART_RE.match(v):
                        own_labels[r] = 'Part ' + PART_RE.match(v).group(1).upper()
            first_own = min(own_labels) if own_labels else None
            for r in range(hr + 1, stop + 1):
                pv = sh.v(r, pc) if pc else None
                if pv and (PART_RE.match(pv) or pv == 'OPTIONAL PIECE'):
                    cur_part = 'Part ' + PART_RE.match(pv).group(1).upper() if PART_RE.match(pv) else 'OPTIONAL PIECE'
                part = cur_part
                if own_labels:
                    if r in own_labels:
                        own_part = own_labels[r]
                    if first_own and r < first_own:
                        prev = chr(ord(own_labels[first_own][-1]) - 1)
                        if cur_part and cur_part.startswith('Part') and cur_part[-1] >= own_labels[first_own][-1]:
                            part = f'Part {prev}'
                        else:
                            part = cur_part
                    elif own_part:
                        part = own_part
                    if cur_part == 'OPTIONAL PIECE':
                        part = cur_part
                slot = sh.v(r, slotc) if slotc else None
                ex = sh.v(r, hc + 2)
                sets = sh.v(r, hc)
                reps0 = sh.v(r, hc + 1)
                rowtxt = [x for x in (sets, reps0, ex, slot) if x]
                if any(x.strip().lower() == 'buy out' for x in rowtxt):
                    part = 'Part E'   # 'Buy out' rows are always the session finisher
                    cand = [x for x in (ex, reps0, sets) if x and x.strip().lower() != 'buy out'
                            and not re.fullmatch(r'[0-9. x]+', x.strip().lower())]
                    if cand:
                        if ex is None or ex.strip().lower() == 'buy out':
                            ex = cand[0]
                            if reps0 == cand[0]:
                                reps0 = None
                    elif ex and ex.strip().lower() == 'buy out':
                        ex = None
                if ex is None:
                    if slot and slot.strip().lower() in ('speed', 'strength'):
                        items.append({'cell': cell(r, slotc), 'row': r, 'name': None, 'day_tag': slot.strip(), 'tag_only': True})
                    continue
                d = {'cell': cell(r, hc + 2), 'row': r, 'sets_raw': sets, 'reps_raw': reps0,
                     'name': ex.strip(), 'part': part, 'slot': slot}
                items.append(d)
            tags = [i['day_tag'] for i in items if i.get('tag_only')]
            items = [i for i in items if not i.get('tag_only')]
            # supersets: a line with no SETS, or with 'Superset' in reps, is paired with the line above
            for k, it in enumerate(items):
                s, rp = it['sets_raw'], (it['reps_raw'] or '')
                if k > 0 and (s is None or 'superset' in rp.lower()) and it['part'] == items[k - 1]['part'] \
                        and not (it['part'] or '').startswith('Part E'):
                    it['paired_with_prev'] = True
            widx = (partcols.index(pc) + 1) if pc in partcols else None
            if items:
                out.append(mk_session('BASEBALL', sh.title, 'athlete', bname, wk, day, items, week_index=widx,
                                      source_range=f"{cell(hr, hc)}:{cell(stop, hc + 2)}",
                                      part_col=(L(pc) if pc else None), slot_col=(L(slotc) if slotc else None),
                                      day_tag=(tags[0] if tags else None)))
    return out

def extract_all():
    wb = openpyxl.load_workbook(RAW)
    S = []
    for f in (ex_ul, ex_fatburn, ex_baseball, ex_tracks, ex_athletic, ex_ppl, ex_boxing, ex_crossfit, ex_conjugate):
        S += f(wb)
    return S, wb

if __name__ == '__main__':
    import collections
    S, wb = extract_all()
    c = collections.Counter((s['program'], s['variant']) for s in S)
    for k, v in c.items():
        print(k, v, sum(len(s['items']) for s in S if (s['program'], s['variant']) == k))
    print(len(S), sum(len(s['items']) for s in S))
